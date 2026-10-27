#!/usr/bin/env python3
"""Generate a neat, professionally aligned resume (PDF + DOCX)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor, Twips
from reportlab.lib.colors import HexColor, black
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch, mm
from reportlab.platypus import (
    Flowable,
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OUT_DIR = Path(__file__).resolve().parent

NAVY = HexColor("#0F2942")
ACCENT = HexColor("#1B4F72")
MUTED = HexColor("#555555")
LINE = HexColor("#C5CDD6")

CRM_BULLETS = [
    "Built Resoline CRM, a full-stack business operations platform using React.js, FastAPI (Python), and MySQL for leads, customers, employees, and role-based access.",
    "Implemented lead and customer management with detailed workspaces, document handling, and AWS S3 storage for attachments and photos.",
    "Developed attendance with GPS punch-in/out, late-punch approval workflows, leave management, and live map-based location tracking for field staff.",
    "Designed a Kanban task board with priorities, due dates, comments, time logs, completion tracking, and file attachments.",
    "Added inventory and stock management with warehouses, purchase receipts, transfers, weighted-average costing, and reorder alerts.",
    "Built vehicle tracking and fuel expense claims with meter-photo verification, distance/fuel calculation, and admin approval flows.",
    "Delivered expenses, payroll/salary, operations workspace, service maintenance, dashboards, and monthly reporting modules.",
    "Secured the app with JWT authentication, permission-based roles, REST APIs, and a responsive Tailwind/Radix UI.",
]


class SectionHeader(Flowable):
    """Bold section title with a clean full-width underline."""

    def __init__(self, text: str, width: float):
        super().__init__()
        self.text = text
        self._width = width
        self.height = 16

    def wrap(self, availWidth, availHeight):
        self._width = availWidth
        return availWidth, self.height

    def draw(self):
        self.canv.setFillColor(NAVY)
        self.canv.setFont("Helvetica-Bold", 11)
        self.canv.drawString(0, 5, self.text.upper())
        self.canv.setStrokeColor(ACCENT)
        self.canv.setLineWidth(1.2)
        self.canv.line(0, 1, self._width, 1)


def _styles():
    return {
        "name": ParagraphStyle(
            "Name",
            fontName="Helvetica-Bold",
            fontSize=20,
            textColor=NAVY,
            alignment=TA_CENTER,
            spaceAfter=2,
            leading=24,
        ),
        "contact": ParagraphStyle(
            "Contact",
            fontName="Helvetica",
            fontSize=9,
            textColor=MUTED,
            alignment=TA_CENTER,
            leading=12,
            spaceAfter=2,
        ),
        "body": ParagraphStyle(
            "Body",
            fontName="Helvetica",
            fontSize=9.5,
            textColor=black,
            alignment=TA_JUSTIFY,
            leading=12.5,
            spaceAfter=2,
        ),
        "role": ParagraphStyle(
            "Role",
            fontName="Helvetica-Bold",
            fontSize=10,
            textColor=black,
            leading=13,
        ),
        "role_right": ParagraphStyle(
            "RoleRight",
            fontName="Helvetica",
            fontSize=9.5,
            textColor=MUTED,
            alignment=TA_RIGHT,
            leading=13,
        ),
        "sub": ParagraphStyle(
            "Sub",
            fontName="Helvetica-Oblique",
            fontSize=9,
            textColor=MUTED,
            leading=12,
        ),
        "sub_right": ParagraphStyle(
            "SubRight",
            fontName="Helvetica-Oblique",
            fontSize=9,
            textColor=MUTED,
            alignment=TA_RIGHT,
            leading=12,
        ),
        "bullet": ParagraphStyle(
            "Bullet",
            fontName="Helvetica",
            fontSize=9.5,
            textColor=black,
            leading=12.2,
            leftIndent=10,
            firstLineIndent=-10,
            spaceBefore=1,
            spaceAfter=1,
        ),
        "proj_title": ParagraphStyle(
            "ProjTitle",
            fontName="Helvetica-Bold",
            fontSize=10,
            textColor=black,
            leading=13,
            spaceBefore=3,
            spaceAfter=1,
        ),
        "skill": ParagraphStyle(
            "Skill",
            fontName="Helvetica",
            fontSize=9.5,
            textColor=black,
            leading=12.5,
            spaceAfter=1,
        ),
        "cert": ParagraphStyle(
            "Cert",
            fontName="Helvetica",
            fontSize=9.5,
            textColor=black,
            leading=12,
            alignment=TA_LEFT,
        ),
    }


def _two_col(left, right, styles, left_w=4.7 * inch, right_w=2.0 * inch):
    """Left-aligned title and right-aligned date/meta on one row."""
    data = [[
        Paragraph(left, styles["role"]),
        Paragraph(right, styles["role_right"]),
    ]]
    t = Table(data, colWidths=[left_w, right_w])
    t.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    return t


def _two_col_sub(left, right, styles, left_w=4.7 * inch, right_w=2.0 * inch):
    data = [[
        Paragraph(left, styles["sub"]),
        Paragraph(right, styles["sub_right"]),
    ]]
    t = Table(data, colWidths=[left_w, right_w])
    t.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )
    return t


def build_pdf(path: Path) -> None:
    page_w, page_h = A4
    left_m = 0.6 * inch
    right_m = 0.6 * inch
    content_w = page_w - left_m - right_m
    left_col = content_w * 0.72
    right_col = content_w * 0.28

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=left_m,
        rightMargin=right_m,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
    )
    s = _styles()
    story = []

    # Header
    story.append(Paragraph("AMIYA SAMAL", s["name"]))
    story.append(
        Paragraph(
            "amiyasamal833@gmail.com &nbsp;&nbsp;•&nbsp;&nbsp; "
            "+91 8144928667 &nbsp;&nbsp;•&nbsp;&nbsp; "
            "Bhubaneswar, Odisha, India",
            s["contact"],
        )
    )
    story.append(
        Paragraph("linkedin.com/in/amiyasamal", s["contact"])
    )
    story.append(Spacer(1, 6))

    # About
    story.append(SectionHeader("About Me", content_w))
    story.append(Spacer(1, 4))
    story.append(
        Paragraph(
            "Motivated and detail-oriented fresher with a strong passion for software development. "
            "Enthusiastic about learning new technologies and applying problem-solving skills to "
            "real-world challenges. Passionate about full-stack development, particularly the MERN "
            "stack, and eager to work on innovative projects that enhance user experience and "
            "efficiency. A quick learner with strong analytical skills and a collaborative mindset.",
            s["body"],
        )
    )
    story.append(Spacer(1, 4))

    # Experience
    story.append(SectionHeader("Work Experience", content_w))
    story.append(Spacer(1, 5))
    story.append(
        _two_col(
            "Freelance Web Developer",
            "01/2024 – Present",
            s,
            left_col,
            right_col,
        )
    )
    story.append(
        _two_col_sub("Self-Employed", "Bhubaneswar, India", s, left_col, right_col)
    )
    story.append(
        Paragraph(
            "•&nbsp;&nbsp;Implemented responsive design techniques for better user experience.",
            s["bullet"],
        )
    )
    story.append(
        Paragraph(
            "•&nbsp;&nbsp;Managed website hosting and maintenance for clients.",
            s["bullet"],
        )
    )
    story.append(Spacer(1, 4))

    # Education
    story.append(SectionHeader("Education", content_w))
    story.append(Spacer(1, 5))
    story.append(
        _two_col(
            "Bachelor of Technology in Computer Science",
            "2022 – 2026",
            s,
            left_col,
            right_col,
        )
    )
    story.append(
        _two_col_sub(
            "Gandhi Engineering College",
            "Bhubaneswar, Odisha",
            s,
            left_col,
            right_col,
        )
    )
    story.append(Spacer(1, 3))
    story.append(
        _two_col(
            "Intermediate in Science (ISC) — 70.33%",
            "2020 – 2022",
            s,
            left_col,
            right_col,
        )
    )
    story.append(
        _two_col_sub(
            "Sadhu Goureswar College",
            "Jajpur, Odisha",
            s,
            left_col,
            right_col,
        )
    )
    story.append(Spacer(1, 3))
    story.append(
        _two_col(
            "High School Certificate — 68.83%",
            "2020",
            s,
            left_col,
            right_col,
        )
    )
    story.append(
        _two_col_sub(
            "Kantabania High School",
            "Jajpur, Odisha",
            s,
            left_col,
            right_col,
        )
    )
    story.append(Spacer(1, 4))

    # Skills
    story.append(SectionHeader("Skills", content_w))
    story.append(Spacer(1, 4))
    skills = [
        ("<b>Programming:</b>&nbsp;&nbsp;C, C++, Python, Java"),
        (
            "<b>Web Development:</b>&nbsp;&nbsp;HTML, CSS, JavaScript, React.js, Node.js, "
            "Express.js, MongoDB, FastAPI"
        ),
        ("<b>Databases:</b>&nbsp;&nbsp;MySQL, SQLite"),
        ("<b>Tools:</b>&nbsp;&nbsp;Git, GitHub, REST APIs, AWS S3"),
        ("<b>Soft Skills:</b>&nbsp;&nbsp;Problem Solving, Communication, Teamwork"),
    ]
    for line in skills:
        story.append(Paragraph(line, s["skill"]))
    story.append(Spacer(1, 4))

    # Projects
    story.append(SectionHeader("Projects", content_w))
    story.append(Spacer(1, 3))

    story.append(Paragraph("Personal Portfolio Website", s["proj_title"]))
    story.append(
        Paragraph(
            "Designed and developed a personal portfolio website to showcase projects and skills. "
            "<i>Technologies: HTML, CSS, JavaScript, Bootstrap</i>",
            s["body"],
        )
    )

    story.append(
        Paragraph(
            "Resoline CRM — Full-Stack Business Operations Platform",
            s["proj_title"],
        )
    )
    for item in CRM_BULLETS:
        story.append(Paragraph(f"•&nbsp;&nbsp;{item}", s["bullet"]))

    story.append(Spacer(1, 4))

    # Certifications
    story.append(SectionHeader("Certifications", content_w))
    story.append(Spacer(1, 4))
    story.append(
        Paragraph(
            "Introduction to Python &nbsp;&nbsp;|&nbsp;&nbsp; "
            "Full-Stack Web Development &nbsp;&nbsp;|&nbsp;&nbsp; "
            "Data Analysis",
            s["cert"],
        )
    )

    doc.build(story)


def _set_run_font(run, name="Calibri", size=10, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def _add_bottom_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1B4F72")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _heading(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    _set_run_font(run, size=11, bold=True, color=RGBColor(0x0F, 0x29, 0x42))
    _add_bottom_border(p)
    return p


def _row(doc, left: str, right: str, left_bold=True, italic=False):
    """Single paragraph with left text and right-aligned tab stop for dates."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r1 = p.add_run(left)
    _set_run_font(
        r1,
        size=10 if left_bold else 9,
        bold=left_bold,
        italic=italic and not left_bold,
        color=RGBColor(0x33, 0x33, 0x33) if italic else RGBColor(0x00, 0x00, 0x00),
    )
    r2 = p.add_run("\t" + right)
    _set_run_font(
        r2,
        size=9.5,
        italic=italic,
        color=RGBColor(0x55, 0x55, 0x55),
    )
    return p


def _bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(text)
    _set_run_font(r, size=9.5)
    return p


def build_docx(path: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # Name
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)
    r = name.add_run("AMIYA SAMAL")
    _set_run_font(r, size=20, bold=True, color=RGBColor(0x0F, 0x29, 0x42))

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(0)
    r = contact.add_run(
        "amiyasamal833@gmail.com  •  +91 8144928667  •  Bhubaneswar, Odisha, India"
    )
    _set_run_font(r, size=9, color=RGBColor(0x55, 0x55, 0x55))

    linkedin = doc.add_paragraph()
    linkedin.alignment = WD_ALIGN_PARAGRAPH.CENTER
    linkedin.paragraph_format.space_after = Pt(4)
    r = linkedin.add_run("linkedin.com/in/amiyasamal")
    _set_run_font(r, size=9, color=RGBColor(0x55, 0x55, 0x55))

    # About
    _heading(doc, "About Me")
    about = doc.add_paragraph()
    about.paragraph_format.space_before = Pt(4)
    about.paragraph_format.space_after = Pt(2)
    r = about.add_run(
        "Motivated and detail-oriented fresher with a strong passion for software development. "
        "Enthusiastic about learning new technologies and applying problem-solving skills to "
        "real-world challenges. Passionate about full-stack development, particularly the MERN "
        "stack, and eager to work on innovative projects that enhance user experience and "
        "efficiency. A quick learner with strong analytical skills and a collaborative mindset."
    )
    _set_run_font(r, size=9.5)

    # Experience
    _heading(doc, "Work Experience")
    _row(doc, "Freelance Web Developer", "01/2024 – Present")
    _row(doc, "Self-Employed", "Bhubaneswar, India", left_bold=False, italic=True)
    _bullet(doc, "Implemented responsive design techniques for better user experience.")
    _bullet(doc, "Managed website hosting and maintenance for clients.")

    # Education
    _heading(doc, "Education")
    _row(doc, "Bachelor of Technology in Computer Science", "2022 – 2026")
    _row(doc, "Gandhi Engineering College", "Bhubaneswar, Odisha", left_bold=False, italic=True)
    _row(doc, "Intermediate in Science (ISC) — 70.33%", "2020 – 2022")
    _row(doc, "Sadhu Goureswar College", "Jajpur, Odisha", left_bold=False, italic=True)
    _row(doc, "High School Certificate — 68.83%", "2020")
    _row(doc, "Kantabania High School", "Jajpur, Odisha", left_bold=False, italic=True)

    # Skills
    _heading(doc, "Skills")
    skills = [
        ("Programming: ", "C, C++, Python, Java"),
        (
            "Web Development: ",
            "HTML, CSS, JavaScript, React.js, Node.js, Express.js, MongoDB, FastAPI",
        ),
        ("Databases: ", "MySQL, SQLite"),
        ("Tools: ", "Git, GitHub, REST APIs, AWS S3"),
        ("Soft Skills: ", "Problem Solving, Communication, Teamwork"),
    ]
    for label, value in skills:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        r1 = p.add_run(label)
        _set_run_font(r1, size=9.5, bold=True)
        r2 = p.add_run(value)
        _set_run_font(r2, size=9.5)

    # Projects
    _heading(doc, "Projects")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Personal Portfolio Website")
    _set_run_font(r, size=10, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(
        "Designed and developed a personal portfolio website to showcase projects and skills. "
        "Technologies: HTML, CSS, JavaScript, Bootstrap"
    )
    _set_run_font(r, size=9.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Resoline CRM — Full-Stack Business Operations Platform")
    _set_run_font(r, size=10, bold=True)
    for item in CRM_BULLETS:
        _bullet(doc, item)

    # Certifications
    _heading(doc, "Certifications")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(
        "Introduction to Python  |  Full-Stack Web Development  |  Data Analysis"
    )
    _set_run_font(r, size=9.5)

    doc.save(path)


def main() -> None:
    pdf_path = OUT_DIR / "AMIYA_Resume.pdf"
    docx_path = OUT_DIR / "AMIYA_Resume.docx"
    build_pdf(pdf_path)
    build_docx(docx_path)
    print(f"Wrote {pdf_path} ({pdf_path.stat().st_size} bytes)")
    print(f"Wrote {docx_path} ({docx_path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
